"""The look of a Word table: the definition, and its one appearance per document.

These tests read the XML rather than the rendered page, because the XML is
where every one of these decisions is either made or silently dropped. A fill
that carries a theme attribute alongside it, a colour with a ``#``, a second
definition of the same style -- none of those raise anywhere, and all of them
change what a reader sees.
"""

from typing import List

import pandas as pd
import pytest
from docx.oxml.ns import qn
from docx.shared import Cm

from mento.reports.table_style import (
    MAX_EIGHTHS,
    MIN_EIGHTHS,
    TableStyle,
    get_table_style,
    points_to_eighths,
    set_table_style,
)
from mento.results import PASS_MARK, FAIL_MARK, TABLE_LOOK, DocumentBuilder

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@pytest.fixture(autouse=True)
def default_table_style() -> None:
    """Leave the package setter as it was found.

    ``set_table_style`` is module state, so a test that sets it would otherwise
    style every table built after it -- including the ones in other modules.
    """
    set_table_style(TableStyle())


@pytest.fixture
def builder() -> DocumentBuilder:
    return DocumentBuilder(title="Test Document")


def _table_frame() -> pd.DataFrame:
    return pd.DataFrame({"Variable": ["b", "h"], "Value": [20, 60]})


def _widths() -> List[Cm]:
    return [Cm(4), Cm(2)]


def _style_elements(builder: DocumentBuilder, style_id: str) -> List[object]:
    """Every ``w:style`` in the document carrying ``style_id``."""
    return [
        element
        for element in builder.doc.styles.element.findall(qn("w:style"))
        if element.get(qn("w:styleId")) == style_id
    ]


# --- Points to eighths ---


def test_points_to_eighths_converts() -> None:
    """1 pt is 8 eighths, which is what ``w:sz`` is counted in."""
    assert points_to_eighths(1.0) == 8
    assert points_to_eighths(0.5) == 4
    assert points_to_eighths(2.25) == 18


def test_points_to_eighths_clamps_below() -> None:
    """A rule thinner than Word draws is written as the thinnest it draws."""
    assert points_to_eighths(0.05) == MIN_EIGHTHS
    assert points_to_eighths(0.25) == MIN_EIGHTHS


def test_points_to_eighths_clamps_above() -> None:
    assert points_to_eighths(50) == MAX_EIGHTHS
    assert points_to_eighths(12) == MAX_EIGHTHS


def test_points_to_eighths_zero_is_no_line() -> None:
    """Zero is not a hairline: it is the absence of a border."""
    assert points_to_eighths(0) == 0
    assert 'w:val="nil"' in TableStyle(inside_pt=0).definition_xml()


def test_points_to_eighths_rejects_negative() -> None:
    with pytest.raises(ValueError, match="cannot be negative"):
        points_to_eighths(-1)


def test_table_style_rejects_negative_thickness() -> None:
    with pytest.raises(ValueError, match="cannot be negative"):
        TableStyle(top_pt=-0.5)


# --- Colours ---


def test_fill_carries_no_theme_attribute() -> None:
    """The fill names a colour and nothing else.

    Word's own ``Light Shading`` declares ``w:fill="C0C0C0"`` and a theme fill
    beside it, resolves the theme first, and so shows a grey nobody asked for.
    A theme attribute reappearing here would make every colour in ``TableStyle``
    a suggestion.
    """
    xml = TableStyle(band_fill="F2F2F2", header_fill="2C6DC4").definition_xml()
    assert "themeFill" not in xml
    assert "themeColor" not in xml
    assert 'w:fill="F2F2F2"' in xml
    assert 'w:fill="2C6DC4"' in xml


@pytest.mark.parametrize("color", ["#F2F2F2", "rojo", "F2F2F", "F2F2F2F", "0x0000FF"])
def test_invalid_color_fails_in_the_constructor(color: str) -> None:
    """Word drops a malformed colour without a word; the constructor does not."""
    with pytest.raises(ValueError, match="is not a colour"):
        TableStyle(band_fill=color)


def test_empty_color_leaves_the_field_unset() -> None:
    """An empty colour is not written, so what is behind it shows through."""
    xml = TableStyle(header_fill="", text_color="").definition_xml()
    assert 'w:type="firstRow"' in xml
    assert "w:shd" not in xml.split('w:type="band1Horz"')[0]
    assert "<w:rPr><w:color" not in xml


def test_lowercase_hex_is_a_colour() -> None:
    assert 'w:fill="f2f2f2"' in TableStyle(band_fill="f2f2f2").definition_xml()


def test_band_size_must_be_a_row_count() -> None:
    with pytest.raises(ValueError, match="band_size"):
        TableStyle(band_size=0)


# --- The definition ---


def test_defaults_are_the_house_look() -> None:
    """The defaults are a decision, not an accident; this is what they are."""
    style = TableStyle()
    assert style.band_fill == "F2F2F2"
    assert style.band_size == 1
    assert style.header_fill == ""
    assert style.header_color == ""
    assert style.header_bold is True
    assert style.text_color == "323232"  # the same grey as the report body
    assert style.border_color == "404040"
    assert (style.top_pt, style.header_rule_pt, style.bottom_pt) == (1.0, 1.0, 1.0)
    assert (style.inside_pt, style.vertical_pt) == (0.0, 0.0)


def test_style_id_is_the_name_reduced_to_letters_and_digits() -> None:
    assert TableStyle().style_id == "MentoTable"
    assert TableStyle(name="Report 2 tables!").style_id == "Report2tables"


def test_name_without_a_letter_or_digit_is_refused() -> None:
    with pytest.raises(ValueError, match="style id"):
        TableStyle(name="---")


def test_definition_carries_the_banding_rule() -> None:
    """The banding is a rule in the document, not paint on today's cells.

    This is the whole reason for a style: a row added in Word afterwards is
    banded because ``band1Horz`` says every other row is.
    """
    xml = TableStyle(band_fill="EEEEEE", band_size=2).definition_xml()
    assert 'w:tblStyleRowBandSize w:val="2"' in xml
    assert 'w:tblStylePr w:type="band1Horz"' in xml
    assert 'w:fill="EEEEEE"' in xml


def test_no_band_fill_means_no_banding_rule() -> None:
    assert "band1Horz" not in TableStyle(band_fill="").definition_xml()


def test_header_bold_lives_in_the_style() -> None:
    bold = TableStyle().definition_xml()
    plain = TableStyle(header_bold=False).definition_xml()
    assert "<w:rPr><w:b/><w:bCs/></w:rPr>" in bold
    assert "<w:b/>" not in plain


def test_thicknesses_reach_the_borders_in_eighths() -> None:
    xml = TableStyle(top_pt=1.5, header_rule_pt=0.75, bottom_pt=2.0, inside_pt=0.5, vertical_pt=0).definition_xml()
    table_borders = xml.split("<w:tblBorders>")[1].split("</w:tblBorders>")[0]
    assert '<w:top w:val="single" w:sz="12"' in table_borders
    assert '<w:bottom w:val="single" w:sz="16"' in table_borders
    assert '<w:insideH w:val="single" w:sz="4"' in table_borders
    assert '<w:insideV w:val="nil"/>' in table_borders
    # The rule under the header is a cell border of the header row: no
    # table-wide setting names that edge.
    header = xml.split('w:type="firstRow"')[1]
    assert '<w:bottom w:val="single" w:sz="6"' in header


# --- Writing the definition into a document ---


def test_definition_is_written_once_for_three_tables(builder: DocumentBuilder) -> None:
    """Forty tables, one definition. That is what a style buys over paint."""
    for _ in range(3):
        builder.add_table(_table_frame(), _widths())

    assert len(builder.doc.tables) == 3
    assert len(_style_elements(builder, "MentoTable")) == 1
    for table in builder.doc.tables:
        assert table._tbl.tblPr.find(qn("w:tblStyle")).get(qn("w:val")) == "MentoTable"


def test_two_styles_in_one_document_do_not_overwrite_each_other(builder: DocumentBuilder) -> None:
    """A second style takes the next free id rather than the first one's."""
    first = builder.table_style_id()
    second = builder.table_style_id(TableStyle(band_fill="F4F7FB", header_fill="2C6DC4"))

    assert first == "MentoTable"
    assert second == "MentoTable2"
    assert len(_style_elements(builder, "MentoTable")) == 1
    assert len(_style_elements(builder, "MentoTable2")) == 1
    definitions = builder.doc.styles.element.xml
    assert 'w:fill="F2F2F2"' in definitions
    assert 'w:fill="F4F7FB"' in definitions
    # And two entries a user can tell apart in Word's gallery.
    assert 'w:val="Mento Table2"' in definitions


def test_a_style_named_after_a_word_built_in_does_not_replace_it(builder: DocumentBuilder) -> None:
    """``Table Grid`` is Word's; a style named for it is written beside it."""
    style_id = builder.table_style_id(TableStyle(name="Table Grid"))

    assert style_id == "TableGrid2"
    assert len(_style_elements(builder, "TableGrid")) == 1


def test_the_package_setter_reaches_a_new_document() -> None:
    """``set_table_style`` is how a report is configured, like the language."""
    set_table_style(TableStyle(band_fill="F4F7FB", header_fill="2C6DC4"))
    assert get_table_style().band_fill == "F4F7FB"

    builder = DocumentBuilder(title="Test Document")
    builder.add_table(_table_frame(), _widths())

    assert 'w:fill="F4F7FB"' in builder.doc.styles.element.xml


def test_a_document_can_override_the_package_setting() -> None:
    builder = DocumentBuilder(title="Test Document", table_style=TableStyle(band_fill="ABCDEF"))
    builder.add_table(_table_frame(), _widths())

    assert 'w:fill="ABCDEF"' in builder.doc.styles.element.xml


def test_set_table_style_refuses_anything_else() -> None:
    with pytest.raises(TypeError, match="TableStyle"):
        set_table_style("MentoTable")  # type: ignore[arg-type]


# --- What the table asks the style for ---


def test_first_column_flag_is_off_and_first_row_on(builder: DocumentBuilder) -> None:
    """The header is a header; the first column is not.

    python-docx creates every table with the first-column flag on, which is
    what made the built-in style bold that column and the builder undo it cell
    by cell. Turning the flag off is the same fix made once.
    """
    builder.add_table(_table_frame(), _widths())
    look = builder.doc.tables[0]._tbl.tblPr.find(qn("w:tblLook"))

    assert look.get(qn("w:firstColumn")) == "0"
    assert look.get(qn("w:lastColumn")) == "0"
    assert look.get(qn("w:firstRow")) == "1"
    assert look.get(qn("w:val")) == "0420"
    for attribute, value in TABLE_LOOK.items():
        assert look.get(qn(f"w:{attribute}")) == value


def test_body_cells_are_not_bolded_by_hand(builder: DocumentBuilder) -> None:
    """No direct bold anywhere: the style decides, so an edit keeps deciding."""
    builder.add_table(_table_frame(), _widths())
    table = builder.doc.tables[0]

    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    assert run.font.bold is None


# --- What the style must not take over ---


def test_verdict_shading_wins_over_the_style(builder: DocumentBuilder) -> None:
    """Green and red are the only place colour means something in the report.

    They are direct cell formatting, which outranks a table style, so the
    banding cannot swallow them -- but this is the one part worth proving
    rather than reasoning about.
    """
    df = pd.DataFrame(
        {
            "Variable": ["Shear", "Flexure", "Units"],
            "Ok?": [PASS_MARK, FAIL_MARK, "kN"],
        }
    )
    builder.add_table_status(df, _widths())
    table = builder.doc.tables[0]

    fills = [table.rows[i + 1].cells[1]._element.tcPr.find(qn("w:shd")) for i in range(3)]
    assert fills[0].get(qn("w:fill")) == "C6EFCE"
    assert fills[1].get(qn("w:fill")) == "FFC7CE"
    # A row that holds no verdict keeps the style's banding, uncoloured.
    assert fills[2] is None
    # And the tables still point at the style that bands them.
    assert table._tbl.tblPr.find(qn("w:tblStyle")).get(qn("w:val")) == "MentoTable"


def test_column_widths_are_still_respected(builder: DocumentBuilder) -> None:
    """The style sets no widths, and must not disturb the ones set here."""
    builder.add_table(_table_frame(), [Cm(4), Cm(2)])
    table = builder.doc.tables[0]

    assert table.autofit is False
    assert table._tbl.tblPr.find(qn("w:tblLayout")).get(qn("w:type")) == "fixed"
    for row in table.rows:
        assert round(row.cells[0].width.cm, 2) == 4.0
        assert round(row.cells[1].width.cm, 2) == 2.0


# --- Cell padding, the field that decides how tall a long table is ---


def test_cell_padding_is_written_in_twentieths() -> None:
    """``w:tblCellMar`` counts in twips, so 1.4 pt is 28 of them."""
    margins = TableStyle(cell_padding_pt=1.4).definition_xml().split("<w:tblCellMar>")[1]
    assert '<w:top w:w="28" w:type="dxa"/>' in margins
    assert '<w:bottom w:w="28" w:type="dxa"/>' in margins
    # The left and right indent is Word's own and does not follow the field:
    # it costs a table nothing in height, which is the only budget that is
    # tight here.
    assert '<w:left w:w="108" w:type="dxa"/>' in margins
    assert '<w:right w:w="108" w:type="dxa"/>' in margins


def test_default_padding_is_tighter_than_words() -> None:
    """The default is what makes a forty-row annex close on one page."""
    assert TableStyle().cell_padding_pt < 1.4


def test_negative_padding_is_refused() -> None:
    with pytest.raises(ValueError, match="cell_padding_pt"):
        TableStyle(cell_padding_pt=-0.5)


def test_a_header_colour_reaches_the_header_row(builder: DocumentBuilder) -> None:
    """``header_color`` is the header's text, and only the header's.

    It is written into the ``firstRow`` part of the definition rather than the
    style's own run properties, so the body keeps ``text_color``.
    """
    style = TableStyle(header_fill="2C6DC4", header_color="FFFFFF")
    header = style.definition_xml().split('w:type="firstRow"')[1]

    assert '<w:color w:val="FFFFFF"/>' in header
    assert '<w:color w:val="FFFFFF"/>' not in style.definition_xml().split("<w:tblPr>")[0]
    assert f'<w:color w:val="{style.text_color}"/>' in style.definition_xml().split("<w:tblPr>")[0]

    builder.add_table(_table_frame(), _widths())
    other_id = builder.table_style_id(style)
    assert 'w:val="FFFFFF"' in builder.doc.styles.element.xml
    assert other_id in builder.doc.styles.element.xml
