"""The single Word report of an element, and where a report is written."""

from __future__ import annotations

import os
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from mento import Concrete_ACI_318_19, Forces, Node, OneWaySlab, RectangularBeam, ShearWall, SteelBar, set_language
from mento import MPa, cm, kN, kNm, m, mm
from mento.reports.views import safe_file_name


def _beam(label: str = "DOC") -> RectangularBeam:
    concrete = Concrete_ACI_318_19(name="C25", f_c=25 * MPa)
    steel = SteelBar(name="ADN420", f_y=420 * MPa)
    beam = RectangularBeam(label=label, concrete=concrete, steel_bar=steel, width=20 * cm, height=50 * cm, c_c=2.5 * cm)
    beam.set_longitudinal_rebar_bot(n1=4, d_b1=16 * mm)
    beam.set_longitudinal_rebar_top(n1=2, d_b1=12 * mm)
    beam.set_transverse_rebar(n_stirrups=1, d_b=8 * mm, s_l=20 * cm)
    return beam


def _checked_node(label: str = "DOC") -> Node:
    node = Node(_beam(label), Forces(label="ELU", M_y=50 * kNm, V_z=80 * kN))
    node.check()
    return node


def _headings(document: object, level: int = 1) -> list[str]:
    return [p.text for p in document.paragraphs if p.style.name == f"Heading {level}"]  # type: ignore[attr-defined]


@pytest.fixture()
def in_tmp_path(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.chdir(tmp_path)
    return tmp_path


# ============================================================================
# One document, flexure and shear
# ============================================================================


def test_one_document_holds_the_flexure_and_the_shear_report() -> None:
    buffer = BytesIO()

    _checked_node().results_detailed_doc(buffer)

    buffer.seek(0)
    document = Document(buffer)
    assert _headings(document) == ["Beam DOC flexure check", "Beam DOC shear check"]
    # Six tables of flexure and six of shear, as the two separate reports hold.
    assert len(document.tables) == 12


def test_the_credit_line_is_written_once() -> None:
    buffer = BytesIO()

    _checked_node().results_detailed_doc(buffer)

    buffer.seek(0)
    credits = [p.text for p in Document(buffer).paragraphs if p.text.startswith("Made with mento")]
    assert len(credits) == 1


def test_a_path_is_written_where_it_says(tmp_path: Path) -> None:
    target = tmp_path / "out" / "memoria.docx"
    target.parent.mkdir()

    _checked_node().results_detailed_doc(target)

    assert _headings(Document(str(target))) == ["Beam DOC flexure check", "Beam DOC shear check"]


def test_no_path_names_the_file_after_the_element(in_tmp_path: Path) -> None:
    _checked_node().results_detailed_doc()

    assert os.listdir(in_tmp_path) == ["Beam DOC check ACI 318-19.docx"]


def test_a_check_that_has_not_run_is_left_out() -> None:
    beam = _beam()
    Node(beam, Forces(label="ELU", V_z=80 * kN)).check_shear()
    buffer = BytesIO()

    beam.results_detailed_doc(buffer)

    buffer.seek(0)
    document = Document(buffer)
    assert _headings(document) == ["Beam DOC shear check"]
    assert any(p.text.startswith("Made with mento") for p in document.paragraphs)


def test_nothing_checked_warns_and_writes_nothing(in_tmp_path: Path) -> None:
    beam = _beam()
    Node(beam, Forces(label="ELU", V_z=80 * kN))

    with pytest.warns(UserWarning, match="No check has been performed"):
        beam.results_detailed_doc()

    assert os.listdir(in_tmp_path) == []


def test_the_single_document_follows_the_report_language() -> None:
    buffer = BytesIO()
    set_language("es")
    try:
        _checked_node().results_detailed_doc(buffer)
    finally:
        set_language("en")

    buffer.seek(0)
    assert _headings(Document(buffer)) == [
        "Verificación a flexión de viga DOC",
        "Verificación a corte de viga DOC",
    ]


def test_a_slab_is_reported_as_a_slab(in_tmp_path: Path) -> None:
    concrete = Concrete_ACI_318_19(name="C25", f_c=25 * MPa)
    steel = SteelBar(name="ADN420", f_y=420 * MPa)
    slab = OneWaySlab(label="L1", concrete=concrete, steel_bar=steel, width=1 * m, height=20 * cm, c_c=2.5 * cm)
    slab.set_slab_longitudinal_rebar_bot(d_b1=12 * mm, s_b1=15 * cm)
    Node(slab, Forces(label="ELU", M_y=30 * kNm, V_z=40 * kN)).check()

    slab.results_detailed_doc()

    assert os.listdir(in_tmp_path) == ["Slab L1 check ACI 318-19.docx"]


def test_a_wall_writes_its_shear_report_to_the_target() -> None:
    concrete = Concrete_ACI_318_19(name="C25", f_c=25 * MPa)
    steel = SteelBar(name="ADN420", f_y=420 * MPa)
    wall = ShearWall(
        label="T1", concrete=concrete, steel_bar=steel, thickness=20 * cm, length=3 * m, height=3 * m, c_c=20 * mm
    )
    wall.set_horizontal_rebar(d_b=10 * mm, s=20 * cm)
    wall.set_vertical_rebar(d_b=10 * mm, s=20 * cm)
    node = Node(wall, Forces(label="ELU", V_z=300 * kN))
    node.check_shear()
    buffer = BytesIO()

    node.results_detailed_doc(buffer)

    buffer.seek(0)
    assert _headings(Document(buffer)) == ["Shear Wall T1 shear check"]


# ============================================================================
# File names
# ============================================================================


@pytest.mark.parametrize("character", list('\\/:*?"<>|'))
def test_every_character_a_file_name_cannot_hold_is_replaced(character: str) -> None:
    assert safe_file_name(f"V1{character}2.docx") == "V1-2.docx"


def test_a_label_with_a_slash_still_saves(in_tmp_path: Path) -> None:
    """``V1/2`` used to ask for a file inside a directory ``V1``: FileNotFoundError."""
    node = _checked_node(label="V1/2")

    node.flexure_results_detailed_doc()
    node.shear_results_detailed_doc()
    node.results_detailed_doc()

    assert sorted(os.listdir(in_tmp_path)) == [
        "Beam V1-2 check ACI 318-19.docx",
        "Beam V1-2 flexure check ACI 318-19.docx",
        "Beam V1-2 shear check ACI 318-19.docx",
    ]
    # The file name is what had to change; the heading keeps the label as written.
    document = Document(str(in_tmp_path / "Beam V1-2 check ACI 318-19.docx"))
    assert _headings(document)[0] == "Beam V1/2 flexure check"
